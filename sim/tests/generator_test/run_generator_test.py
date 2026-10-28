#!/usr/bin/env python3
"""Tests of CSR_Generation.py itself (TEST-001 and TEST-002 of the specification).

Runs the tool on the reference workbook and on bad inputs and checks files, ports, exit status and determinism.
Usage : python run_generator_test.py [WORKBOOK]   (default: CSR_GOLDEN_FORMAT.xlsx, or CSR_Example.xlsx)
Needs : Python 3 and openpyxl. Exit 0 when every check passes.
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..'))
TOOL = os.path.join(ROOT, 'CSR_Generation.py')
results = []


def check(name, ok, detail=''):
    results.append(ok)
    print('%s  %s %s' % ('PASS' if ok else 'FAIL', name, detail))


def run(args, cwd=None):
    return subprocess.run([sys.executable, TOOL] + args, capture_output=True, text=True, cwd=cwd)


def ports_of(sv):
    text = open(sv, encoding='utf-8').read()
    m = re.search(r'\)\s*\(\s*(.*?)\);', text, re.S)
    return re.findall(r'(?m)^\s*(input|output)\s+(?:wire|logic)\s*(?:\[[^\]]+\])?\s*(\w+)\s*,?\s*$', m.group(1))


def expected_field_ports(ws):
    """Field ports expected by REQ-IF-003 from the register table of a sheet."""
    exp, reg, started = set(), None, False
    for row in ws.iter_rows(values_only=True):
        a = row[0]
        if not started:
            if a and 'REGISTER' in str(a):
                started = True
            continue
        if a:
            reg = a
        elif row[2] and row[2] != 'reserved':
            t, n = row[4], '%s_%s' % (reg, row[2])
            if t == 'ro':
                exp.add('i_' + n)
            else:
                exp.add('o_' + n)
                if t in ('rwi', 'w1c'):
                    exp |= {'i_hw_wdata_' + n, 'i_hw_we_' + n}
    return exp


def main():
    wb_path = sys.argv[1] if len(sys.argv) > 1 else next(
        p for p in (os.path.join(ROOT, 'CSR_GOLDEN_FORMAT.xlsx'), os.path.join(ROOT, 'CSR_Example.xlsx')) if os.path.exists(p))
    tmp = tempfile.mkdtemp()
    wb = openpyxl.load_workbook(wb_path)
    for sheet, asyn in (('CSR', 0), ('ASYNC_CSR', 1)):
        if sheet not in wb.sheetnames:
            continue
        out = os.path.join(tmp, sheet)
        r = run([wb_path, sheet, '--out', out, '--project', 'Test', '--author', 'Tester (user)'])
        cfg = {str(wb[sheet].cell(i, 1).value): wb[sheet].cell(i, 2).value for i in range(3, 9)}
        mod = cfg['Module Name']
        sv = os.path.join(out, 'RTL', mod + '.sv')
        check('%s: exit 0' % sheet, r.returncode == 0)
        check('%s: module file, file list, debug dump' % sheet,
              os.path.isfile(sv) and os.path.isfile(os.path.join(out, 'RTL', 'filelist.f')) and os.path.isfile(os.path.join(out, 'Debug', 'debug.json')))
        check('%s: synchronizer model only when Asynchronous = 1' % sheet,
              os.path.isfile(os.path.join(out, 'RTL', 'models', 'm_csr_sync.sv')) == bool(asyn))
        head = open(sv, encoding='utf-8').read()[:700]
        check('%s: header fields' % sheet, all(re.search(r'%s\s*:' % k, head) for k in ('Project', 'Author', 'AI Flow', 'AI Flow Page', 'Date', 'Description'))
              and 'Author       : Tester (user)' in head and 'Design Rule v2' in head)
        got = {n for d, n in ports_of(sv)}
        fixed = {'i_clk_apb', 'i_rst_n_apb', 'i_apb_psel', 'i_apb_penable', 'i_apb_pwrite', 'i_apb_paddr', 'i_apb_pwdata', 'i_apb_pstrb',
                 'i_apb_pprot', 'o_apb_pready', 'o_apb_pslverr', 'o_apb_prdata', 'i_protect_en', 'i_slverr_en'}
        if asyn:
            fixed |= {'i_clk_reg', 'i_rst_n_reg'}
        check('%s: ports = fixed ports + field ports' % sheet, got == fixed | expected_field_ports(wb[sheet]), '(%d ports)' % len(got))
        # TEST-002: same inputs, same files except the Date field
        out2 = os.path.join(tmp, sheet + '_2')
        run([wb_path, sheet, '--out', out2, '--project', 'Test', '--author', 'Tester (user)'])
        a = re.sub(r'Date\s*:.*', '', open(sv, encoding='utf-8').read())
        b = re.sub(r'Date\s*:.*', '', open(os.path.join(out2, 'RTL', mod + '.sv'), encoding='utf-8').read())
        check('%s: repeatable output' % sheet, a == b)
    # negative inputs
    r = run([wb_path, 'NO_SUCH_SHEET', '--out', os.path.join(tmp, 'x')])
    check('missing sheet: exit status 1', r.returncode == 1)
    bad = os.path.join(tmp, 'bad_protocol.xlsx')
    w2 = openpyxl.load_workbook(wb_path)
    sheet0 = 'CSR' if 'CSR' in w2.sheetnames else w2.sheetnames[0]
    for i in range(1, 12):
        if w2[sheet0].cell(i, 1).value == 'Protocol':
            w2[sheet0].cell(i, 2).value = 'AXI'
    w2.save(bad)
    r = run([bad, sheet0, '--out', os.path.join(tmp, 'y')])
    check('protocol other than APB: exit status 1 and a message', r.returncode == 1 and 'not implemented' in r.stdout)
    miss = os.path.join(tmp, 'missing_cfg.xlsx')
    w3 = openpyxl.load_workbook(wb_path)
    for i in range(1, 12):
        if w3[sheet0].cell(i, 1).value == 'Data Width':
            w3[sheet0].cell(i, 1).value = 'Width of data'
    w3.save(miss)
    r = run([miss, sheet0, '--out', os.path.join(tmp, 'z')])
    check('missing configuration value: exit status 1 and a message naming it', r.returncode == 1 and 'data_wd' in r.stdout)
    # an empty row inside the register table is skipped
    w4 = openpyxl.load_workbook(wb_path)
    ws = w4[sheet0]
    start = next(i for i in range(1, 40) if ws.cell(i, 1).value and 'REGISTER' in str(ws.cell(i, 1).value))
    ws.insert_rows(start + 3)
    gap = os.path.join(tmp, 'gap.xlsx')
    w4.save(gap)
    r = run([gap, sheet0, '--out', os.path.join(tmp, 'g')])
    check('empty row in the register table is skipped', r.returncode == 0)
    shutil.rmtree(tmp, ignore_errors=True)
    print('RESULT: %s (%d of %d checks)' % ('PASS' if all(results) else 'FAIL', sum(results), len(results)))
    return 0 if all(results) else 1


if __name__ == '__main__':
    sys.exit(main())
