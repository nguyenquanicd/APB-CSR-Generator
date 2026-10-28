#--------------------------------------------------------------------
# Project: APB-CSR-Generator
# Author: Trthinh (Ethan)
# Function: CSR Generation Tool
# Page: VLSI Technology
#
# Version 2.1: a hardware write to a w1c field has priority over a software clear in the same cycle
#   (version 2.0: software first); rwi keeps software first.
# Version 2.0: the generated RTL follows the VLSIT RTL Design Rule v2
#   (spec_template/VLSIT_RTL_Design_Rule.md of VLSIT_RTL_Generator_AI_Model):
#   names (i_clk_apb, i_apb_*, LPARA_*, reg_*, w_*), port layout, header,
#   process labels, case defaults, no latch, no unused declaration.
#   The behavior and the timing of version 1.2 are unchanged.
#--------------------------------------------------------------------

import argparse
import getpass
import os
import shutil
import sys
from datetime import date, datetime

VERSION = "2.1"
RSVD_KEY = 'reserved'
FLOW_NAME = "VLSIT RTL Generator AI Automation Flow"
FLOW_PAGE = "https://github.com/nguyenquanicd/APB-CSR-Generator"
SYNC_MODULE = "m_csr_sync"


def bit_width(field):
    if ":" in field:
        temp = field.replace("[", "")
        temp = temp.replace("]", "")
        temp = temp.split(':')
        return int(temp[0]) - int(temp[1]) + 1
    else:
        return 1


def range_text(width):
    """Packed range of a field port: empty for one bit (rule TYP-06)."""
    return f"[{width - 1}:0]" if width > 1 else ""


class port:
    """One port or internal signal: direction input / output / logic, width in bits, optional width text."""
    def __init__(self, name, direction, width, width_text=None):
        if width < 1:
            print(f"ERROR: The width of port {name} is incorrect")
            sys.exit(1)
        self.name = name
        self.direction = direction
        self.width = width
        self.width_text = width_text if width_text is not None else range_text(width)

    def logic_rtl(self):
        space = f" {self.width_text}" if self.width_text else ""
        return f"  logic{space} {self.name};\n"


class port_group:
    """Ports of one function (rule FIL-06): a title comment and the ports."""
    def __init__(self, title, ports):
        self.title = title
        self.ports = ports


def ports_rtl(groups):
    """Port list in four aligned columns: direction, type, width, name (rules FIL-05, FIL-06, TYP-02)."""
    rows = []
    for g in groups:
        for p in g.ports:
            typ = "wire" if p.direction == "input" else "logic"
            rows.append((p.direction, typ, p.width_text, p.name))
    dir_col = max(len(r[0]) for r in rows) + 1
    typ_col = max(len(r[1]) for r in rows) + 1
    wid_col = max(len(r[2]) for r in rows)
    wid_col = wid_col + 1 if wid_col else 0
    total = len(rows)
    out = ""
    count = 0
    for gi, g in enumerate(groups):
        if gi:
            out += "\n"
        out += f"  // {g.title}\n"
        for p in g.ports:
            typ = "wire" if p.direction == "input" else "logic"
            count += 1
            comma = "," if count < total else ""
            out += "  " + p.direction.ljust(dir_col) + typ.ljust(typ_col) + p.width_text.ljust(wid_col) + p.name + comma + "\n"
    return out


class register:
    def __init__(self, name, offset, reset, description, bit_list, clock, reset_n):
        self.name = name
        self.offset = offset
        self.reset = reset
        self.description = description
        self.bit = bit_list
        self.clock = clock
        self.reset_n = reset_n
        self.port_list = []
        self.logic_list = []
        self.has_we = any(b.type != 'ro' for b in self.bit)
        self.lcparam = f"  localparam logic [31:0] LPARA_{self.name.upper()}_OFFSET = 32'h{self.offset[2:]},\n"
        self.we_rtl = ''
        self.rtl = ''
        self.gen_port_logic()
        self.gen_rtl()

    def gen_port_logic(self):
        n = self.name
        if self.has_we:
            self.we_rtl = (f"  assign w_we_{n} = w_wr_en & (reg_addr == LPARA_{n.upper()}_OFFSET[LPARA_ADDR_W-1:0]);\n")
            self.logic_list.append(port(f"w_we_{n}", "logic", 1))
        for b in self.bit:
            if b.type == 'ro' and b.name != RSVD_KEY:
                self.port_list.append(port(f"i_{n}_{b.name}", "input", b.width))
            elif b.type == 'rw':
                self.port_list.append(port(f"o_{n}_{b.name}", "output", b.width))
                self.logic_list.append(port(f"reg_{n}_{b.name}", "logic", b.width))
                self.logic_list.append(port(f"w_{n}_{b.name}_nxt", "logic", b.width))
            elif b.type == 'rwi' or b.type == 'w1c':
                self.port_list.append(port(f"o_{n}_{b.name}", "output", b.width))
                self.port_list.append(port(f"i_hw_wdata_{n}_{b.name}", "input", b.width))
                self.port_list.append(port(f"i_hw_we_{n}_{b.name}", "input", 1))
                self.logic_list.append(port(f"reg_{n}_{b.name}", "logic", b.width))
                self.logic_list.append(port(f"w_{n}_{b.name}_nxt", "logic", b.width))
        self.logic_list.append(port(f"w_{n}_value", "logic", 32, "[LPARA_DATA_W-1:0]"))

    def gen_rtl(self):
        n = self.name
        flops = [b for b in self.bit if b.type != 'ro']
        temp = ""
        for b in self.bit:
            b.next_value_cal()
            self.rtl += b.rtl
            if b.type == 'ro':
                if b.name != RSVD_KEY:
                    temp += f"    w_{n}_value{b.field} = i_{n}_{b.name};\n"
            else:
                temp += f"    w_{n}_value{b.field} = reg_{n}_{b.name};\n"
        if flops:
            reset_part = ""
            update_part = ""
            for b in flops:
                reset_val = b.reset_val.replace('0x', "")
                reset_part += f"      reg_{n}_{b.name} <= {b.width}'h{reset_val};\n"
                update_part += f"      reg_{n}_{b.name} <= w_{n}_{b.name}_nxt;\n"
            self.rtl += f'''
  // Register {n}: flops of the writable fields, reset to the field reset values.
  always_ff @(posedge {self.clock}, negedge {self.reset_n}) begin : p_ff_{n}
    if (!{self.reset_n}) begin
{reset_part[:-1]}
    end else begin
{update_part[:-1]}
    end
  end

  // Register {n}: output ports follow the flops.
'''
            for b in flops:
                self.rtl += f"  assign o_{n}_{b.name} = reg_{n}_{b.name};\n"
        self.rtl += f'''
  // Register {n}: read value from the fields; reserved bits read as 0.
  always_comb begin : p_com_{n}_value
    w_{n}_value = '0;
{temp[:-1]}
  end
'''


class bit:
    def __init__(self, name, field, type, reset_val, description, reg_group, async_option):
        self.name = name
        self.field = field
        self.type = type
        self.reset_val = reset_val
        self.description = description
        self.port_list = []
        self.logic_list = []
        self.width = bit_width(self.field)
        self.reg_group = reg_group
        self.rtl = ''

    def next_value_cal(self):
        r = self.reg_group
        nxt = f"w_{r}_{self.name}_nxt"
        cur = f"reg_{r}_{self.name}"
        space = " " * (len(f"  assign {nxt} = ("))
        if self.type == 'rw':
            self.rtl += f'''
  // {r}_{self.name} next value (rw): a software write loads it, otherwise hold.
  assign {nxt} = (w_we_{r}) ? reg_wdata{self.field} : {cur};
'''
        elif self.type == 'rwi':
            self.rtl += f'''
  // {r}_{self.name} next value (rwi): software write, then hardware write, then hold.
  assign {nxt} = (w_we_{r}) ? reg_wdata{self.field}
{space}: (i_hw_we_{r}_{self.name}) ? i_hw_wdata_{r}_{self.name}
{space}: {cur};
'''
        elif self.type == 'w1c':
            self.rtl += f'''
  // {r}_{self.name} next value (w1c): hardware write, then software 1 clears, then hold.
  assign {nxt} = (i_hw_we_{r}_{self.name}) ? i_hw_wdata_{r}_{self.name}
{space}: (w_we_{r}) ? (~reg_wdata{self.field} & {cur})
{space}: {cur};
'''


class apb_protocol:
    def __init__(self, data_wd, addr_wd, async_option):
        self.data_wd = data_wd
        self.addr_wd = addr_wd
        self.async_option = int(async_option)
        self.clk_bus = "i_clk_apb"
        self.rst_bus = "i_rst_n_apb"
        self.clk_reg = "i_clk_reg" if self.async_option else "i_clk_apb"
        self.rst_reg = "i_rst_n_reg" if self.async_option else "i_rst_n_apb"
        self.port_groups = []
        self.logic_list = []
        self.gen_port()
        self.gen_logic()
        self.apb_access()
        self.rw_rtl()

    def gen_port(self):
        g = [port_group("APB clock and reset", [port(self.clk_bus, "input", 1), port(self.rst_bus, "input", 1)])]
        if self.async_option:
            g.append(port_group("Register clock and reset", [port("i_clk_reg", "input", 1), port("i_rst_n_reg", "input", 1)]))
        g.append(port_group("APB slave interface", [
            port("i_apb_psel", "input", 1),
            port("i_apb_penable", "input", 1),
            port("i_apb_pwrite", "input", 1),
            port("i_apb_paddr", "input", self.addr_wd, "[LPARA_ADDR_W-1:0]"),
            port("i_apb_pwdata", "input", self.data_wd, "[LPARA_DATA_W-1:0]"),
            port("i_apb_pstrb", "input", 4, "[LPARA_STRB_W-1:0]"),
            port("i_apb_pprot", "input", 3),
            port("o_apb_pready", "output", 1),
            port("o_apb_pslverr", "output", 1),
            port("o_apb_prdata", "output", self.data_wd, "[LPARA_DATA_W-1:0]")]))
        g.append(port_group("Access check control", [port("i_protect_en", "input", 1), port("i_slverr_en", "input", 1)]))
        self.port_groups = g

    def gen_logic(self):
        L = self.logic_list
        for name in ("w_apb_setup", "w_apb_prot_ok", "w_apb_slverr", "w_apb_complete", "w_apb_wr", "w_apb_rd"):
            L.append(port(name, "logic", 1))
        L.append(port("reg_addr", "logic", self.addr_wd, "[LPARA_ADDR_W-1:0]"))
        L.append(port("reg_wdata", "logic", self.data_wd, "[LPARA_DATA_W-1:0]"))
        L.append(port("reg_slverr", "logic", 1))
        L.append(port("w_rd_en", "logic", 1))
        L.append(port("w_wr_en", "logic", 1))
        L.append(port("w_rdata_nxt", "logic", self.data_wd, "[LPARA_DATA_W-1:0]"))
        L.append(port("reg_pready", "logic", 1))
        L.append(port("reg_rdata", "logic", self.data_wd, "[LPARA_DATA_W-1:0]"))
        L.append(port("w_pready_nxt", "logic", 1))
        if self.async_option:
            for name in ("w_rd_ack", "w_rd_ack_bus", "reg_rd_req", "w_rd_req_nxt", "w_rd_ack_fall", "reg_rd_ack_bus_dly",
                         "w_wr_ack", "w_wr_ack_bus", "reg_wr_req", "w_wr_req_nxt", "w_wr_ack_fall", "reg_wr_ack_bus_dly",
                         "w_rd_req_reg", "reg_rd_req_reg_dly", "w_wr_req_reg", "reg_wr_req_reg_dly"):
                L.append(port(name, "logic", 1))
            L.append(port("reg_rdata_bus", "logic", self.data_wd, "[LPARA_DATA_W-1:0]"))
        else:
            L.append(port("reg_rd_capture", "logic", 1))
            L.append(port("reg_wr_capture", "logic", 1))

    def apb_access(self):
        cb, rb = self.clk_bus, self.rst_bus
        self.access = f"""
  // APB decode: the SETUP phase is PSEL high and PENABLE low and lasts one cycle per transfer.
  // A transfer is in error when the protection, the word alignment or the byte strobes fail
  // while i_slverr_en is 1.
  assign w_apb_setup    = i_apb_psel & ~i_apb_penable;
  assign w_apb_prot_ok  = i_protect_en ? ~i_apb_pprot[1] : 1'b1;
  assign w_apb_slverr   = i_slverr_en ? (~w_apb_prot_ok
                                         | (|i_apb_paddr[1:0])
                                         | (~&i_apb_pstrb & i_apb_pwrite)) : 1'b0;

  // Transfer without error, split into a write and a read.
  assign w_apb_complete = w_apb_setup & ~w_apb_slverr;
  assign w_apb_wr       = i_apb_pwrite & w_apb_complete;
  assign w_apb_rd       = ~i_apb_pwrite & w_apb_complete;

  // Address capture: holds the address of the last transfer without error.
  always_ff @(posedge {cb}, negedge {rb}) begin : p_ff_addr
    if (!{rb})
      reg_addr <= '0;
    else if (w_apb_complete)
      reg_addr <= i_apb_paddr;
  end

  // Write data capture: holds the data of the last write without error.
  always_ff @(posedge {cb}, negedge {rb}) begin : p_ff_wdata
    if (!{rb})
      reg_wdata <= '0;
    else if (w_apb_wr)
      reg_wdata <= i_apb_pwdata;
  end

  // Error response flop: loaded in every SETUP phase while i_slverr_en is 1, and held otherwise.
  always_ff @(posedge {cb}, negedge {rb}) begin : p_ff_slverr
    if (!{rb})
      reg_slverr <= 1'b0;
    else if (w_apb_setup & i_slverr_en)
      reg_slverr <= w_apb_slverr;
  end

  // PSLVERR output follows the error response flop.
  assign o_apb_pslverr = reg_slverr;
"""

    def rw_rtl(self):
        cb, rb, cr, rr = self.clk_bus, self.rst_bus, self.clk_reg, self.rst_reg
        if self.async_option:
            self.rtl_rw = f'''
  // Read request and acknowledge toggles cross between the APB clock and the register clock.
  // Bus clock, read: the acknowledge is synchronized into the APB clock.
  {SYNC_MODULE} #(
    .PARA_LEVELS (LPARA_SYNC_STAGES)
  ) u_rd_ack_sync (
    .i_clk_dst   ({cb}),
    .i_rst_n_dst ({rb}),
    .i_d         (w_rd_ack),
    .o_q         (w_rd_ack_bus)
  );

  // Bus clock, read: the request is set by a read and cleared by the acknowledge.
  assign w_rd_req_nxt  = w_rd_ack_bus ? 1'b0 : (w_apb_rd ? 1'b1 : reg_rd_req);
  assign w_rd_ack_fall = reg_rd_ack_bus_dly & ~w_rd_ack_bus;

  // Bus clock, read: request flop and delayed acknowledge flop.
  always_ff @(posedge {cb}, negedge {rb}) begin : p_ff_rd_req
    if (!{rb}) begin
      reg_rd_req         <= 1'b0;
      reg_rd_ack_bus_dly <= 1'b0;
    end else begin
      reg_rd_req         <= w_rd_req_nxt;
      reg_rd_ack_bus_dly <= w_rd_ack_bus;
    end
  end

  // Bus clock, write: the acknowledge is synchronized into the APB clock.
  {SYNC_MODULE} #(
    .PARA_LEVELS (LPARA_SYNC_STAGES)
  ) u_wr_ack_sync (
    .i_clk_dst   ({cb}),
    .i_rst_n_dst ({rb}),
    .i_d         (w_wr_ack),
    .o_q         (w_wr_ack_bus)
  );

  // Bus clock, write: the request is set by a write and cleared by the acknowledge.
  assign w_wr_req_nxt  = w_wr_ack_bus ? 1'b0 : (w_apb_wr ? 1'b1 : reg_wr_req);
  assign w_wr_ack_fall = reg_wr_ack_bus_dly & ~w_wr_ack_bus;

  // Bus clock, write: request flop and delayed acknowledge flop.
  always_ff @(posedge {cb}, negedge {rb}) begin : p_ff_wr_req
    if (!{rb}) begin
      reg_wr_req         <= 1'b0;
      reg_wr_ack_bus_dly <= 1'b0;
    end else begin
      reg_wr_req         <= w_wr_req_nxt;
      reg_wr_ack_bus_dly <= w_wr_ack_bus;
    end
  end

  // Register clock, read: the request is synchronized into the register clock.
  {SYNC_MODULE} #(
    .PARA_LEVELS (LPARA_SYNC_STAGES)
  ) u_rd_req_sync (
    .i_clk_dst   ({cr}),
    .i_rst_n_dst ({rr}),
    .i_d         (reg_rd_req),
    .o_q         (w_rd_req_reg)
  );

  // Register clock, read: a new request is a rising edge; the acknowledge is the delayed request.
  always_ff @(posedge {cr}, negedge {rr}) begin : p_ff_rd_req_reg
    if (!{rr})
      reg_rd_req_reg_dly <= 1'b0;
    else
      reg_rd_req_reg_dly <= w_rd_req_reg;
  end

  // Read enable: pulse on the rising edge of the request; the acknowledge is the delayed request.
  assign w_rd_en  = ~reg_rd_req_reg_dly & w_rd_req_reg;
  assign w_rd_ack = reg_rd_req_reg_dly;

  // Register clock, write: the request is synchronized into the register clock.
  {SYNC_MODULE} #(
    .PARA_LEVELS (LPARA_SYNC_STAGES)
  ) u_wr_req_sync (
    .i_clk_dst   ({cr}),
    .i_rst_n_dst ({rr}),
    .i_d         (reg_wr_req),
    .o_q         (w_wr_req_reg)
  );

  // Register clock, write: a new request is a rising edge; the acknowledge is the delayed request.
  always_ff @(posedge {cr}, negedge {rr}) begin : p_ff_wr_req_reg
    if (!{rr})
      reg_wr_req_reg_dly <= 1'b0;
    else
      reg_wr_req_reg_dly <= w_wr_req_reg;
  end

  // Write enable: pulse on the rising edge of the request; the acknowledge is the delayed request.
  assign w_wr_en  = ~reg_wr_req_reg_dly & w_wr_req_reg;
  assign w_wr_ack = reg_wr_req_reg_dly;

  // PREADY: one pulse when a read or write is acknowledged, or after an error SETUP phase.
  assign w_pready_nxt = w_rd_ack_fall | w_wr_ack_fall | (w_apb_slverr & w_apb_setup);

  // PREADY flop: loaded with the next value, cleared by the reset.
  always_ff @(posedge {cb}, negedge {rb}) begin : p_ff_pready
    if (!{rb})
      reg_pready <= 1'b0;
    else
      reg_pready <= w_pready_nxt;
  end

  // PREADY output follows the PREADY flop.
  assign o_apb_pready = reg_pready;
'''
        else:
            self.rtl_rw = f'''
  // Read and write enables: the transfer without error, delayed by one cycle.
  always_ff @(posedge {cb}, negedge {rb}) begin : p_ff_capture
    if (!{rb}) begin
      reg_rd_capture <= 1'b0;
      reg_wr_capture <= 1'b0;
    end else begin
      reg_rd_capture <= w_apb_rd;
      reg_wr_capture <= w_apb_wr;
    end
  end

  // Enables follow the capture flops.
  assign w_rd_en = reg_rd_capture;
  assign w_wr_en = reg_wr_capture;

  // PREADY next value: one cycle after the read or write enable, or after an error SETUP phase.
  assign w_pready_nxt = w_wr_en | w_rd_en | (w_apb_slverr & w_apb_setup);

  // PREADY flop: loaded with the next value, cleared by the reset.
  always_ff @(posedge {cb}, negedge {rb}) begin : p_ff_pready
    if (!{rb})
      reg_pready <= 1'b0;
    else
      reg_pready <= w_pready_nxt;
  end

  // PREADY output follows the PREADY flop.
  assign o_apb_pready = reg_pready;
'''

    def prdata_rtl(self, reg_list):
        cb, rb, cr, rr = self.clk_bus, self.rst_bus, self.clk_reg, self.rst_reg
        items = ""
        for r in reg_list:
            items += f"      LPARA_{r.name.upper()}_OFFSET[LPARA_ADDR_W-1:0]: w_rdata_nxt = w_{r.name}_value;\n"
        rtl = f'''
  // Read data mux: value of the register at the captured address, 0 for any other address.
  always_comb begin : p_com_rdata
    w_rdata_nxt = '0;
    case (reg_addr)
{items}      default: w_rdata_nxt = '0;
    endcase
  end
'''
        if self.async_option:
            rtl += f'''
  // Read data snapshot in the register clock domain, loaded by the read enable of that domain.
  always_ff @(posedge {cr}, negedge {rr}) begin : p_ff_rdata
    if (!{rr})
      reg_rdata <= '0;
    else if (w_rd_en)
      reg_rdata <= w_rdata_nxt;
  end

  // Read data resynchronized to the APB clock: reg_rdata is stable when the acknowledge
  // pulse arrives, so that pulse samples it without a bit synchronizer.
  always_ff @(posedge {cb}, negedge {rb}) begin : p_ff_rdata_bus
    if (!{rb})
      reg_rdata_bus <= '0;
    else if (w_rd_ack_fall)
      reg_rdata_bus <= reg_rdata;
  end

  // PRDATA output follows the read data flop of the APB clock.
  assign o_apb_prdata = reg_rdata_bus;
'''
        else:
            rtl += f'''
  // Read data register, loaded by the read enable.
  always_ff @(posedge {cb}, negedge {rb}) begin : p_ff_rdata
    if (!{rb})
      reg_rdata <= '0;
    else if (w_rd_en)
      reg_rdata <= w_rdata_nxt;
  end

  // PRDATA output follows the read data register.
  assign o_apb_prdata = reg_rdata;
'''
        self.rtl_prdata = rtl


class sfr:
    def __init__(self, module_name, addr_wd, data_wd, reg_list, protocol, async_option, sheet="", project="APB-CSR-Generator", author=None):
        self.module_name = module_name
        self.addr_wd = addr_wd
        self.data_wd = data_wd
        self.async_option = async_option
        self.sheet = sheet
        self.project = project
        self.author = author if author else f"{getpass.getuser()} (user)"
        if protocol.lower() == 'apb':
            self.protocol = apb_protocol(data_wd=data_wd, addr_wd=addr_wd, async_option=async_option)
        else:
            print(f"ERROR: protocol {protocol} is not implemented")
            sys.exit(1)
        self.rtl = ''
        self.reg_list = reg_list
        self.gen_port_logic_rtl()

    def gen_port_logic_rtl(self):
        groups = list(self.protocol.port_groups)
        for r in self.reg_list:
            if r.port_list:
                groups.append(port_group(f"Register {r.name}", r.port_list))
        self.port_rtl = ports_rtl(groups)
        self.logic_rtl = ''
        for l in self.protocol.logic_list:
            self.logic_rtl += l.logic_rtl()
        for r in self.reg_list:
            for lr in r.logic_list:
                self.logic_rtl += lr.logic_rtl()

    def gen_json(self):
        reg_json = ''
        bracket_o = "{"
        bracket_c = "}"
        for r in self.reg_list:
            bit_json = ''
            reg_json += f'''        {bracket_o}
        "Register Name": "{r.name}",
        "Offset": "{r.offset}",
        "Reset Value": "{r.reset}",
        "Descriprtion": "{r.description}",
        "Bits": ['''
            for b in r.bit:
                bit_json += f'''
            {bracket_o}
                "Bit Name": "{b.name}",
                "Bit Type": "{b.type}",
                "Bit Field": "{b.field}",
                "Reset Value": "{b.reset_val}",
                "Description": "{b.description}"
            {bracket_c},'''
            reg_json += bit_json[:-1] + f"\n{' ' * 10}]\n" + f"{' ' * 8}{bracket_c},\n"

        reg_json = reg_json[:-2]
        self.json = f'''    {bracket_o}
    "Project": "APB-CSR-Generator",
    "Module": "JSON debug file",
    "Function": "Used to check the configuration after generation",
    "Author": "Trthinh (Ethan)",
    "Generator": "{getpass.getuser()}",
    "Generator version": "{VERSION}",
    "Page": "VLSI Technology",
    "Table - Configuration": [
    {bracket_o}
        "Module Name": "{self.module_name}",
        "Address Width": "{self.addr_wd}",
        "Data Width": "{self.data_wd}",
        "Protocol": "APB",
        "Asynchornous": "{self.async_option}"
    {bracket_c}
    ],
    "Table - Register Definition": [
{reg_json}
    ]
{bracket_c}'''

    def gen_rtl(self):
        async_text = "two clocks, i_clk_apb and i_clk_reg" if int(self.async_option) else "one clock, i_clk_apb"
        header = f'''/*
 * Project      : {self.project}
 * Author       : {self.author}
 * AI Flow      : {FLOW_NAME}
 * AI Flow Page : {FLOW_PAGE}
 * Date         : {date.today().isoformat()}
 * Description  : Control and status registers with an APB4 slave interface.
 *                Clocks: {async_text}.
 *                Generated by APB-CSR-Generator v{VERSION}; do not edit by hand.
 *                Source: sheet {self.sheet} of the CSR workbook.
 *                The RTL follows the VLSIT RTL Design Rule v2.
 */
`timescale 1ns/1ps
'''
        localparam_offset = ("  localparam logic [31:0] LPARA_ADDR_W = " + f"32'd{self.addr_wd},\n"
                             "  localparam logic [31:0] LPARA_DATA_W = " + f"32'd{self.data_wd},\n"
                             "  localparam logic [31:0] LPARA_STRB_W = " + f"32'd{self.data_wd // 8},\n")
        if int(self.async_option):
            localparam_offset += "  localparam logic [31:0] LPARA_SYNC_STAGES = 32'd2,\n"
        reg_rtl = ''
        reg_we = ''
        for r in self.reg_list:
            localparam_offset += r.lcparam
            reg_rtl += r.rtl
            reg_we += r.we_rtl
        localparam_offset = localparam_offset[:-2] + "\n"
        self.protocol.prdata_rtl(self.reg_list)
        we_part = f"\n  // Register write enables: a write enable of the captured address.\n{reg_we}" if reg_we else ""
        body = f'''
module {self.module_name} #(
{localparam_offset}) (
{self.port_rtl});

  // Internal signal declarations
{self.logic_rtl}{self.protocol.access}{we_part}{self.protocol.rtl_rw}{reg_rtl}{self.protocol.rtl_prdata}
endmodule
'''
        self.rtl = header + body


SYNC_MODEL = f'''/*
 * Project      : APB-CSR-Generator
 * Author       : Trthinh (Ethan) (user)
 * AI Flow      : {FLOW_NAME}
 * AI Flow Page : {FLOW_PAGE}
 * Date         : 2026-10-28
 * Description  : Behavioral synchronizer model of the generated CSR.
 *                One bit, PARA_LEVELS flip-flops in series.
 *                Placeholder: replace this file by a wrapper of the synchronizer cell of the
 *                technology library before tape-out; keep the module name and the port names.
 *                The RTL follows the VLSIT RTL Design Rule v2.
 */
`timescale 1ns/1ps

module {SYNC_MODULE} #(
  // Parameter Description: number of flip-flop stages of the synchronizer.
  // Parameter values: 2 or more; default is 2.
  // Parameter unit: count
  parameter logic [31:0] PARA_LEVELS = 2
) (
  // Destination clock and reset
  input  wire  i_clk_dst,
  input  wire  i_rst_n_dst,

  // Synchronized bit
  input  wire  i_d,
  output logic o_q
);

  // Internal signal declarations
  logic [PARA_LEVELS-1:0] reg_sync;

  // Synchronizer chain: the input is shifted through PARA_LEVELS flip-flops, cleared by the reset.
  always_ff @(posedge i_clk_dst, negedge i_rst_n_dst) begin : p_ff_sync
    if (!i_rst_n_dst)
      reg_sync <= '0;
    else
      reg_sync <= {{reg_sync[PARA_LEVELS-2:0], i_d}};
  end

  // Output mapping: the last stage is the synchronized bit.
  assign o_q = reg_sync[PARA_LEVELS-1];

endmodule
'''


def read_workbook(path, sheet_name):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet_name]
    start_config = 0
    start_reg = 0
    reg_list = []
    cfg = {}
    start_col = 0
    reg_col = start_col
    offset_col = start_col + 1
    reset_col = start_col + 5
    des_col = start_col + 6
    bit_name_col = start_col + 2
    bit_field_col = start_col + 3
    bit_type_col = start_col + 4
    reg_name = ''
    raw_regs = []
    for row in range(1, ws.max_row + 1):
        if start_reg == 0:
            if start_config == 0:
                if 'Table - Configuration' in str(ws[row][start_col].value):
                    start_config = 1
            else:
                label = str(ws[row][start_col].value)
                if 'Module Name' in label:
                    cfg['module_name'] = ws[row][start_col + 1].value
                if 'Protocol' in label:
                    cfg['protocol'] = ws[row][start_col + 1].value
                if 'Data Width' in label:
                    cfg['data_wd'] = int(ws[row][start_col + 1].value)
                if 'Address Width' in label:
                    cfg['addr_wd'] = int(ws[row][start_col + 1].value)
                if 'Asynchronous' in label:
                    cfg['async_option'] = int(ws[row][start_col + 1].value)
                if 'REGISTER' in label:
                    start_reg = 1
        else:
            if ws[row][0].value is not None:
                reg_name = ws[row][reg_col].value
                raw_regs.append({"name": reg_name, "offset": ws[row][offset_col].value, "reset": ws[row][reset_col].value,
                                 "description": ws[row][des_col].value, "bits": []})
            else:
                if not raw_regs:
                    continue
                bit_name = ws[row][bit_name_col].value
                bit_field = ws[row][bit_field_col].value
                if bit_name is None and bit_field is None:
                    continue    # an empty row of the sheet is not a field
                raw_regs[-1]["bits"].append((bit_name, bit_field, ws[row][bit_type_col].value,
                                             ws[row][reset_col].value, ws[row][des_col].value))
    return cfg, raw_regs


def main():
    parser = argparse.ArgumentParser(description="APB-CSR-Generator: APB4 CSR RTL from an Excel register map")
    parser.add_argument("workbook")
    parser.add_argument("sheet")
    parser.add_argument("--out", default=os.getcwd(), help="folder that receives RTL/ and Debug/ (default: current folder)")
    parser.add_argument("--project", default="APB-CSR-Generator", help="Project field of the file header")
    parser.add_argument("--author", default=None, help="Author field of the file header, '<name> (user)'")
    args = parser.parse_args()

    print("Start the CSR Generation!")
    print("The register configured will be shown below!")
    script_path = os.path.dirname(os.path.realpath(__file__))
    cfg, raw_regs = read_workbook(args.workbook, args.sheet)
    for key in ('module_name', 'protocol', 'data_wd', 'addr_wd', 'async_option'):
        if key not in cfg:
            print(f"ERROR: the configuration table of the sheet {args.sheet} has no value for {key}")
            sys.exit(1)
    reg_list = []
    clock = 'i_clk_reg' if int(cfg['async_option']) else 'i_clk_apb'
    reset_n = 'i_rst_n_reg' if int(cfg['async_option']) else 'i_rst_n_apb'
    for rr in raw_regs:
        bit_list = []
        for (bit_name, bit_field, bit_type, bit_reset, bit_des) in rr["bits"]:
            print("Register: ", rr["name"], ", Bit name: ", bit_name, ", Bit Type: ", bit_type)
            bit_list.append(bit(name=bit_name, field=bit_field, type=bit_type, reset_val=bit_reset, description=bit_des,
                                reg_group=rr["name"], async_option=cfg['async_option']))
        reg_list.append(register(name=rr["name"], offset=rr["offset"], reset=rr["reset"], description=rr["description"],
                                 bit_list=bit_list, clock=clock, reset_n=reset_n))
    s = sfr(cfg['module_name'], cfg['addr_wd'], cfg['data_wd'], reg_list, cfg['protocol'], cfg['async_option'],
            sheet=args.sheet, project=args.project, author=args.author)
    s.gen_rtl()

    out = os.path.abspath(args.out)
    print("The output path: ", out)
    print("The tool path: ", script_path)
    rtl_path = out + "/RTL"
    os.makedirs(rtl_path, exist_ok=True)
    with open(rtl_path + '/' + s.module_name + '.sv', 'w') as f:
        f.write(s.rtl)
    models_path = rtl_path + '/models'
    os.makedirs(models_path, exist_ok=True)
    if int(cfg['async_option']):
        # an existing model is kept: it may be the synchronizer cell wrapper that replaced the placeholder
        if os.path.exists(models_path + '/' + SYNC_MODULE + '.sv'):
            print(f"Keep the existing models/{SYNC_MODULE}.sv")
        else:
            with open(models_path + '/' + SYNC_MODULE + '.sv', 'w') as f:
                f.write(SYNC_MODEL)
    with open(rtl_path + '/filelist.f', 'w') as f:
        f.write('// filelist.f - generated by APB-CSR-Generator; paths are relative to the folder of this file.\n')
        if int(cfg['async_option']):
            f.write(f'models/{SYNC_MODULE}.sv\n')
        f.write(s.module_name + '.sv\n')

    s.gen_json()
    debug_path = out + "/Debug"
    os.makedirs(debug_path, exist_ok=True)
    with open(debug_path + '/debug.json', 'w') as f:
        f.write(s.json)
    print("End of the CSR generation! Thank you.")


if __name__ == "__main__":
    main()
